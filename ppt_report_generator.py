import pandas as pd
import os

from pptx import Presentation
from pptx.util import Inches, Pt


# ==================================================
# LOAD DATA
# ==================================================

sentiment_df = pd.read_csv(
    "data/aspect_sentiment_results.csv"
)

emotion_df = pd.read_csv(
    "data/emotion_results.csv"
)

customer_df = pd.read_csv(
    "data/customer_segments.csv"
)

alert_df = pd.read_csv(
    "data/sentiment_alerts.csv"
)

recommendation_df = pd.read_csv(
    "data/business_recommendations.csv"
)


# ==================================================
# CREATE REPORT FOLDER
# ==================================================

os.makedirs(
    "reports",
    exist_ok=True
)


# ==================================================
# CREATE PRESENTATION
# ==================================================

prs = Presentation()

prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


# ==================================================
# HELPER FUNCTIONS
# ==================================================

def add_title(slide, title):

    title_box = slide.shapes.add_textbox(
        Inches(0.6),
        Inches(0.3),
        Inches(12),
        Inches(0.7)
    )

    text_frame = title_box.text_frame

    paragraph = text_frame.paragraphs[0]

    paragraph.text = title

    paragraph.font.size = Pt(28)
    paragraph.font.bold = True


def add_bullets(
    slide,
    items,
    left=0.8,
    top=1.5,
    width=11.5,
    height=5
):

    box = slide.shapes.add_textbox(
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height)
    )

    text_frame = box.text_frame

    text_frame.word_wrap = True

    for i, item in enumerate(items):

        if i == 0:
            paragraph = text_frame.paragraphs[0]
        else:
            paragraph = text_frame.add_paragraph()

        paragraph.text = item

        paragraph.font.size = Pt(20)

        paragraph.space_after = Pt(12)


def add_table(
    slide,
    data,
    left=0.6,
    top=1.5,
    width=12,
    height=4.8
):

    rows = len(data)

    cols = len(data[0])

    table_shape = slide.shapes.add_table(
        rows,
        cols,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height)
    )

    table = table_shape.table

    for r in range(rows):

        for c in range(cols):

            cell = table.cell(
                r,
                c
            )

            cell.text = str(
                data[r][c]
            )

            for paragraph in cell.text_frame.paragraphs:

                paragraph.font.size = Pt(14)

                if r == 0:
                    paragraph.font.bold = True


# ==================================================
# SLIDE 1 — TITLE
# ==================================================

slide = prs.slides.add_slide(
    prs.slide_layouts[6]
)

add_title(
    slide,
    "Multi-Aspect Sentiment Analysis System"
)

add_bullets(
    slide,
    [
        "Amazon Product Review Analysis",
        "",
        "Sentiment Analysis • Aspect Extraction • Emotion Analysis",
        "Customer Segmentation • Alerts • Business Intelligence"
    ],
    top=2.2
)


# ==================================================
# SLIDE 2 — PROJECT OBJECTIVE
# ==================================================

slide = prs.slides.add_slide(
    prs.slide_layouts[6]
)

add_title(
    slide,
    "Project Objective"
)

add_bullets(
    slide,
    [
        "Analyze customer reviews at aspect level.",
        "Identify positive, neutral and negative sentiment.",
        "Track sentiment trends over time.",
        "Classify customer emotion patterns.",
        "Generate actionable business recommendations.",
        "Provide an interactive dashboard and real-time API."
    ]
)


# ==================================================
# SLIDE 3 — DATASET
# ==================================================

slide = prs.slides.add_slide(
    prs.slide_layouts[6]
)

add_title(
    slide,
    "Dataset Summary"
)

total_reviews = len(
    sentiment_df
)

customers = customer_df[
    "reviewerID"
].nunique()

data = [
    ["Metric", "Value"],
    ["Aspect-level Records", total_reviews],
    ["Customers", customers],
    ["Products", 1],
    ["Review Period", "2012 - 2014"]
]

add_table(
    slide,
    data
)


# ==================================================
# SLIDE 4 — SENTIMENT ANALYSIS
# ==================================================

slide = prs.slides.add_slide(
    prs.slide_layouts[6]
)

add_title(
    slide,
    "Overall Sentiment Analysis"
)

sentiment_counts = (
    sentiment_df["sentiment"]
    .value_counts()
)

data = [
    ["Sentiment", "Count"]
]

for sentiment, count in sentiment_counts.items():

    data.append([
        sentiment,
        count
    ])

add_table(
    slide,
    data
)


# ==================================================
# SLIDE 5 — ASPECT ANALYSIS
# ==================================================

slide = prs.slides.add_slide(
    prs.slide_layouts[6]
)

add_title(
    slide,
    "Aspect-Based Sentiment Analysis"
)

aspect_table = pd.crosstab(
    sentiment_df["aspect"],
    sentiment_df["sentiment"]
)

for sentiment in [
    "Positive",
    "Neutral",
    "Negative"
]:

    if sentiment not in aspect_table.columns:

        aspect_table[sentiment] = 0

aspect_table = aspect_table[
    [
        "Positive",
        "Neutral",
        "Negative"
    ]
]

data = [
    [
        "Aspect",
        "Positive",
        "Neutral",
        "Negative"
    ]
]

for aspect, row in aspect_table.iterrows():

    data.append([
        aspect,
        row["Positive"],
        row["Neutral"],
        row["Negative"]
    ])

add_table(
    slide,
    data,
    height=5.2
)


# ==================================================
# SLIDE 6 — EMOTION ANALYSIS
# ==================================================

slide = prs.slides.add_slide(
    prs.slide_layouts[6]
)

add_title(
    slide,
    "8-Emotion Classification"
)

emotion_counts = (
    emotion_df["emotion"]
    .value_counts()
)

data = [
    ["Emotion", "Count"]
]

for emotion, count in emotion_counts.items():

    data.append([
        emotion,
        count
    ])

add_table(
    slide,
    data
)


# ==================================================
# SLIDE 7 — CUSTOMER SEGMENTATION
# ==================================================

slide = prs.slides.add_slide(
    prs.slide_layouts[6]
)

add_title(
    slide,
    "Customer Segmentation"
)

segment_counts = (
    customer_df["Segment"]
    .value_counts()
)

data = [
    ["Customer Segment", "Customers"]
]

for segment, count in segment_counts.items():

    data.append([
        segment,
        count
    ])

add_table(
    slide,
    data
)


# ==================================================
# SLIDE 8 — NEGATIVE SENTIMENT ALERTS
# ==================================================

slide = prs.slides.add_slide(
    prs.slide_layouts[6]
)

add_title(
    slide,
    "Negative Sentiment Alerts"
)

alerts = alert_df[
    alert_df["Alert"] != "Normal"
]

alert_items = [
    f"Negative sentiment spike months detected: {len(alerts)}"
]

if len(alerts) > 0:

    for _, row in alerts.iterrows():

        alert_items.append(
            f"{row['Month']} - "
            f"{row['Negative_Percentage']:.2f}% negative"
        )

else:

    alert_items.append(
        "No negative sentiment spikes detected."
    )

add_bullets(
    slide,
    alert_items
)


# ==================================================
# SLIDE 9 — BUSINESS RECOMMENDATIONS
# ==================================================

slide = prs.slides.add_slide(
    prs.slide_layouts[6]
)

add_title(
    slide,
    "Business Recommendations"
)

recommendation_items = []

# Detect aspect column

if "Aspect" in recommendation_df.columns:

    aspect_column = "Aspect"

elif "aspect" in recommendation_df.columns:

    aspect_column = "aspect"

elif "Unnamed: 0" in recommendation_df.columns:

    aspect_column = "Unnamed: 0"

else:

    aspect_column = None


# Detect recommendation column

if "Recommendation" in recommendation_df.columns:

    recommendation_column = "Recommendation"

elif "recommendation" in recommendation_df.columns:

    recommendation_column = "recommendation"

else:

    recommendation_column = None


if (
    aspect_column is not None
    and recommendation_column is not None
):

    for _, row in recommendation_df.iterrows():

        recommendation_items.append(
            f"{row[aspect_column]}: "
            f"{row[recommendation_column]}"
        )

else:

    recommendation_items.append(
        "Business recommendations generated "
        "for analyzed product aspects."
    )


# Keep slide readable

recommendation_items = recommendation_items[:8]

add_bullets(
    slide,
    recommendation_items,
    top=1.3,
    height=5.8
)


# ==================================================
# SLIDE 10 — SYSTEM ARCHITECTURE
# ==================================================

slide = prs.slides.add_slide(
    prs.slide_layouts[6]
)

add_title(
    slide,
    "System Architecture"
)

add_bullets(
    slide,
    [
        "Amazon Review Dataset",
        "↓",
        "Data Preprocessing",
        "↓",
        "Aspect Extraction",
        "↓",
        "Sentiment + Emotion Analysis",
        "↓",
        "Trend Analysis + Customer Segmentation",
        "↓",
        "Dashboard + API + Reports + Alerts"
    ],
    top=1.2
)


# ==================================================
# SLIDE 11 — TECHNOLOGIES
# ==================================================

slide = prs.slides.add_slide(
    prs.slide_layouts[6]
)

add_title(
    slide,
    "Technologies Used"
)

add_bullets(
    slide,
    [
        "Python",
        "Pandas",
        "NLTK / TextBlob",
        "Scikit-learn based analytics workflow",
        "Streamlit",
        "Flask",
        "ReportLab",
        "python-pptx",
        "Matplotlib"
    ]
)


# ==================================================
# SLIDE 12 — LIMITATIONS
# ==================================================

slide = prs.slides.add_slide(
    prs.slide_layouts[6]
)

add_title(
    slide,
    "Project Limitations"
)

add_bullets(
    slide,
    [
        "The current dataset contains reviews for one product.",
        "Therefore, genuine competitor comparison requires multi-product data.",
        "Customer membership levels are derived from review activity.",
        "Emotion classification uses keyword-based lexical patterns.",
        "Multilingual and sarcasm detection require additional modules."
    ]
)


# ==================================================
# SLIDE 13 — CONCLUSION
# ==================================================

slide = prs.slides.add_slide(
    prs.slide_layouts[6]
)

add_title(
    slide,
    "Conclusion"
)

add_bullets(
    slide,
    [
        "Multi-aspect sentiment analysis was implemented.",
        "Sentiment trends and negative sentiment alerts were generated.",
        "Eight-emotion classification was implemented.",
        "Customer sentiment patterns were segmented.",
        "Business recommendations were generated.",
        "Interactive dashboard and real-time API were implemented.",
        "Automated PDF and PPT reporting are supported."
    ]
)


# ==================================================
# SAVE PPT
# ==================================================

ppt_path = (
    "reports/sentiment_analysis_report.pptx"
)

prs.save(
    ppt_path
)

print(
    "PowerPoint report generated successfully!"
)

print(
    "Saved at:",
    ppt_path
)